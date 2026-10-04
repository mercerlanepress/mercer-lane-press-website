"""Customer-package verification and independent LibreOffice render/recalculation in CI."""
import json, os, subprocess, zipfile, pathlib, tempfile, copy, re, xml.etree.ElementTree as ET
import openpyxl
from pypdf import PdfReader
ROOT=pathlib.Path(__file__).resolve().parent.parent
OUT=ROOT/'free-review-qa';OUT.mkdir(exist_ok=True)
with zipfile.ZipFile(ROOT/'public/downloads/Mercer_Lane_Free_Performance_Review_Starter_Kit.zip') as z:
    assert len([n for n in z.namelist() if not n.endswith('/')])==6
    assert all(n.endswith(('.pdf','.docx','.xlsx')) for n in z.namelist())
    z.extractall(OUT)
KIT=OUT/'Mercer_Lane_Free_Performance_Review_Starter_Kit'
XLSX=KIT/'Mercer_Lane_Free_Performance_Review_Template.xlsx'
w=openpyxl.load_workbook(XLSX)
assert len(w.sheetnames)==5
for s in w:
    assert s.protection.sheet and not s.protection.password
    for row in s:
        for c in row:
            if c.data_type=='f':
                assert c.protection.locked
                assert len(c.value)<8192
                # OOXML schema validation and LibreOffice accept longer literals;
                # Excel's formula parser rejects any quoted string over 255 characters.
                for literal in re.findall(r'"(?:[^"]|"")*"', c.value):
                    assert len(literal[1:-1].replace('""','"'))<=255, (s.title,c.coordinate,'Excel formula literal too long')
assert not w['PERFORMANCE REVIEW']['D10'].protection.locked
assert not w['GOALS & ACTIONS']['A9'].protection.locked
assert not w['AI REVIEW PROMPT']['D10'].protection.locked
assert len(w['AI REVIEW PROMPT'].data_validations.dataValidation)==9
for s in w:
    for dv in s.data_validations.dataValidation:assert dv.showErrorMessage
with zipfile.ZipFile(XLSX) as z:
    assert not any('externalLink' in n or 'vbaProject' in n for n in z.namelist())
for f in KIT.glob('*.pdf'):
    r=PdfReader(f);assert not r.trailer['/Root'].get('/OpenAction');assert '/JavaScript' not in str(r.trailer)
    assert not r.trailer['/Root'].get('/Names',{}).get('/EmbeddedFiles')
    assert all(p.extract_text().strip() for p in r.pages)
    if '_A4' in f.name:assert abs(float(r.pages[0].mediabox.width)-595.276)<1
    if '_US_Letter' in f.name:assert abs(float(r.pages[0].mediabox.width)-612)<1
# Fixtures are QA-only, separate from the customer workbook. openpyxl is used for checks, not authorship.
spec=json.loads((ROOT/'src/data/performance-review-prompt.json').read_text())
rows=[10,12,14,17,20,23,26]
vals=['Marketing Coordinator','Q3 2026','Delivered September campaign.','Clear project updates.','Improve handover documentation.','Manager reviews checklist.','Use checklist before next launch.']
for i,(kind,desc) in enumerate(spec['outputs'].items()):
    sample=openpyxl.load_workbook(XLSX);s=sample['AI REVIEW PROMPT']
    for row,val in zip(rows,vals):s[f'D{row}']=val
    s['D30']=kind;s['D29']='Supportive';p=OUT/f'sample-{i}.xlsx';sample.save(p)
    subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/mlp-review-qa','--headless','--convert-to','xlsx','--outdir',str(OUT/'recalculated'),str(p)],check=True)
    v=openpyxl.load_workbook(OUT/'recalculated'/p.name,data_only=True)['AI REVIEW PROMPT']['A34'].value
    assert isinstance(v,str) and desc in v and 'Tone: Supportive' in v and all(x in v for x in vals),(kind,str(v)[:100])
    (OUT/f'prompt-{i}.txt').write_text(v)
subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/mlp-review-qa','--headless','--convert-to','pdf','--outdir',str(OUT),str(KIT/'Employee_Performance_Review_Form.docx')],check=True)
docpdf=OUT/'Employee_Performance_Review_Form.pdf';assert len(PdfReader(docpdf).pages)==3
subprocess.run(['pdftoppm','-r','110','-png',str(docpdf),str(OUT/'word-page')],check=True)
print('Customer files validated; five output types recalculated independently; Word rendered as three pages.')
